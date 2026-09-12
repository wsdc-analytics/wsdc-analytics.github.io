#!/usr/bin/env python3
"""Build draft WSDC Registry Event Competitor List export workbook (xlsx).

Audience: scoring vendors (machine export), not manual organizer entry.
Open in Excel or upload to Google Sheets.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "WSDC_Registry_Event_Competitor_List_v0.1.xlsx"

# Enums for Lookups sheet (column order = named ranges via A1 refs)
ROLES = ["Leader", "Follower"]
DIVISIONS = [
    "Newcomer",
    "Novice",
    "Intermediate",
    "Advanced",
    "All-Stars",
    "Champions",
    "Sophisticated",
    "Masters",
    "Juniors",
]
CONTEST_TYPES = ["Jack & Jill", "Strictly", "Other"]
ENTRY_STATUSES = ["registered", "scratched", "no_show", "danced"]
YES_NO = ["Y", "N"]
ID_ABSENT_REASONS = ["first_time_no_id", "unknown", "not_provided", ""]
COUNTRIES_SAMPLE = [
    "USA",
    "Canada",
    "France",
    "Germany",
    "United Kingdom",
    "Italy",
    "Spain",
    "Netherlands",
    "Sweden",
    "Poland",
    "Australia",
    "Japan",
    "South Korea",
    "Brazil",
    "Other",
]

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14, color="1E3A5F")
SECTION_FONT = Font(bold=True, name="Calibri", size=11)
NOTE_FONT = Font(italic=True, name="Calibri", size=10, color="555555")
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
SOFT = PatternFill("solid", fgColor="F8FAFC")
REQ = PatternFill("solid", fgColor="FEF3C7")


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def _autosize(ws, min_w: int = 12, max_w: int = 36) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value is None:
                continue
            length = max(length, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = max(min_w, length)


def _write_lookups(ws) -> dict[str, str]:
    """Return map of enum name -> sheet formula range like Lookups!$A$2:$A$3."""
    ws["A1"] = "role"
    ws["B1"] = "division"
    ws["C1"] = "contest_type"
    ws["D1"] = "entry_status"
    ws["E1"] = "yes_no"
    ws["F1"] = "id_absent_reason"
    ws["G1"] = "event_country_sample"
    _style_header(ws, 1, 7)

    tables = {
        "role": ROLES,
        "division": DIVISIONS,
        "contest_type": CONTEST_TYPES,
        "entry_status": ENTRY_STATUSES,
        "yes_no": YES_NO,
        "id_absent_reason": [r for r in ID_ABSENT_REASONS if r],
        "event_country_sample": COUNTRIES_SAMPLE,
    }
    col_map = {
        "role": 1,
        "division": 2,
        "contest_type": 3,
        "entry_status": 4,
        "yes_no": 5,
        "id_absent_reason": 6,
        "event_country_sample": 7,
    }
    ranges: dict[str, str] = {}
    for name, values in tables.items():
        c = col_map[name]
        for i, v in enumerate(values, start=2):
            ws.cell(row=i, column=c, value=v)
        end = 1 + len(values)
        letter = get_column_letter(c)
        ranges[name] = f"Lookups!${letter}$2:${letter}${end}"
    _autosize(ws)
    ws.freeze_panes = "A2"
    return ranges


def _add_list_validation(ws, ranges: dict[str, str], col_letter: str, enum_key: str, max_row: int = 2005) -> None:
    dv = DataValidation(
        type="list",
        formula1=f"={ranges[enum_key]}",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Invalid value",
        error="Pick a value from the list (Lookups sheet).",
    )
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}6:{col_letter}{max_row}")


def _write_readme(ws) -> None:
    ws["A1"] = "WSDC Registry Event — Competitor List export (DRAFT v0.1)"
    ws["A1"].font = TITLE_FONT
    lines = [
        "",
        "Purpose",
        "Replace / succeed the Score Report «Competitors List» sheet as a machine-export format from scoring systems into WSDC retention + analytics (path toward C1).",
        "Not designed for organizers to type hundreds of names by hand.",
        "",
        "Who fills this",
        "Scoring vendors / systems: export after the event (or after registration close + final updates) in this schema.",
        "Event organizers: only confirm Event_Header values if the scoring system cannot resolve event_id / catalog name.",
        "",
        "How to use",
        "1. Fill Event_Header (one row per event edition).",
        "2. Export Competitors: one row = one person × one role × one division × one contest_type.",
        "3. Dual / second role = second row (same person, different role), not two roles in one cell.",
        "4. Upload xlsx to Google Sheets or submit xlsx as-is. Keep column headers unchanged.",
        "",
        "Identity rules (critical for join to Points Registry / C1)",
        "• wsdc_competitor_id: required when the dancer already has a WSDC Competitor ID; leave blank if none yet.",
        "• Competitor ID exists only after a first awarded point (Rules §3.2.2c) — blank ID is valid Entry.",
        "• has_wsdc_id must be Y or N. If N, set id_absent_reason.",
        "• Do not invent IDs. Names stay in WSDC-controlled storage (PII); public analytics use counts / joined IDs only.",
        "",
        "Join keys (analytics)",
        "• Event: prefer event_id from WSDC catalog; fallback (event_name_canonical, event_year).",
        "• Dancer: wsdc_competitor_id ↔ Points Registry dancer_id when present.",
        "",
        "Sheets in this workbook",
        "• README — this note",
        "• Event_Header — edition metadata (dropdowns where marked)",
        "• Competitors — entry rows (dropdowns on role / division / contest_type / status / flags)",
        "• Lookups — controlled vocabularies (do not rename columns used by dropdowns)",
        "• Vendor_Mapping — checklist for scoring companies",
        "",
        "Status: DRAFT for committee / ops discussion. Not an official WSDC form until adopted.",
    ]
    for i, line in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=line)
        if line in {
            "Purpose",
            "Who fills this",
            "How to use",
            "Identity rules (critical for join to Points Registry / C1)",
            "Join keys (analytics)",
            "Sheets in this workbook",
        }:
            ws.cell(row=i, column=1).font = SECTION_FONT
        elif line.startswith("Status:"):
            ws.cell(row=i, column=1).font = NOTE_FONT
    ws.column_dimensions["A"].width = 110


def _write_event_header(ws, ranges: dict[str, str]) -> None:
    ws["A1"] = "Event_Header — one row per registry event edition"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Yellow columns = required for analytics ingest. Dropdowns where listed."
    ws["A2"].font = NOTE_FONT

    headers = [
        ("event_id", "Required when catalog ID exists; else leave blank and fill canonical name"),
        ("event_name_canonical", "Exact catalog / WSDC listing name — no free-text nicknames"),
        ("event_year", "WSDC results year for this edition (integer)"),
        ("event_start_date", "ISO date YYYY-MM-DD"),
        ("event_end_date", "ISO date YYYY-MM-DD"),
        ("event_country", "Dropdown sample list; extend Lookups as needed"),
        ("event_state_or_region", "State / region if applicable; else blank"),
        ("scoring_system", "Vendor / product name that produced the export"),
        ("score_report_export_version", "Use WSDC_CL_v0.1 for this draft"),
        ("exported_at", "ISO datetime UTC when file was generated"),
        ("exported_by", "System or operator contact"),
    ]
    for col, (h, _) in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)
    _style_header(ws, 4, len(headers))
    for col, (_, hint) in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col, value=hint)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    sample = [
        "",  # event_id TBD
        "Example Swing Event",
        2026,
        "2026-03-14",
        "2026-03-16",
        "USA",
        "CA",
        "ExampleScoringCo",
        "WSDC_CL_v0.1",
        "2026-03-17T12:00:00Z",
        "export@example-scoring.example",
    ]
    for col, val in enumerate(sample, start=1):
        cell = ws.cell(row=6, column=col, value=val)
        cell.border = THIN
        if col in {2, 3, 4, 5, 6, 8, 9, 10}:
            cell.fill = REQ

    # country dropdown on row 6
    dv = DataValidation(
        type="list",
        formula1=f"={ranges['event_country_sample']}",
        allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(dv)
    dv.add("F6:F20")

    ws.row_dimensions[5].height = 48
    _autosize(ws, min_w=14, max_w=28)
    ws.freeze_panes = "A6"


def _write_competitors(ws, ranges: dict[str, str]) -> None:
    ws["A1"] = "Competitors — scoring-system export (one row per person × role × division × contest)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        "Organizers should not type this sheet. Scoring systems map registration/start lists → these columns. "
        "Blank wsdc_competitor_id is allowed when has_wsdc_id=N."
    )
    ws["A2"].font = NOTE_FONT

    headers = [
        "row_id",
        "wsdc_competitor_id",
        "registered_name",
        "preferred_name",
        "role",
        "division",
        "contest_type",
        "entry_status",
        "has_wsdc_id",
        "id_absent_reason",
        "notes_internal",
    ]
    req_cols = {1, 5, 6, 7, 8, 9}  # row_id, role, division, contest, status, has_wsdc_id

    for col, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)
    _style_header(ws, 4, len(headers))

    hints = [
        "Stable unique id from scoring system",
        "WSDC Competitor ID if assigned; else blank",
        "Legal / registration name (PII — WSDC store only)",
        "Optional display name",
        "Dropdown",
        "Dropdown",
        "Dropdown",
        "Dropdown",
        "Y / N",
        "Required if has_wsdc_id=N",
        "Optional vendor note; not for public analytics",
    ]
    for col, hint in enumerate(hints, start=1):
        cell = ws.cell(row=5, column=col, value=hint)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True)

    samples = [
        ["SRC-0001", "12345", "Alex Example", "Alex", "Leader", "Novice", "Jack & Jill", "danced", "Y", "", ""],
        ["SRC-0002", "", "Sam Firsttime", "Sam", "Follower", "Newcomer", "Jack & Jill", "danced", "N", "first_time_no_id", ""],
        ["SRC-0003", "12345", "Alex Example", "Alex", "Follower", "Novice", "Jack & Jill", "registered", "Y", "", "dual role → second row"],
        ["SRC-0004", "67890", "Jordan Strictly", "", "Leader", "Intermediate", "Strictly", "scratched", "Y", "", ""],
    ]
    for r, row in enumerate(samples, start=6):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = THIN
            if c in req_cols:
                cell.fill = REQ

    # Empty template rows with light borders for demo paste
    for r in range(10, 16):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c, value="")
            cell.border = THIN
            if c in req_cols:
                cell.fill = SOFT

    _add_list_validation(ws, ranges, "E", "role")
    _add_list_validation(ws, ranges, "F", "division")
    _add_list_validation(ws, ranges, "G", "contest_type")
    _add_list_validation(ws, ranges, "H", "entry_status")
    _add_list_validation(ws, ranges, "I", "yes_no")
    _add_list_validation(ws, ranges, "J", "id_absent_reason")

    ws.row_dimensions[5].height = 40
    _autosize(ws, min_w=12, max_w=22)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A4:K15"


def _write_vendor_mapping(ws) -> None:
    ws["A1"] = "Vendor_Mapping — what scoring systems should map"
    ws["A1"].font = TITLE_FONT
    headers = ["Source in scoring system (typical)", "Target column", "Rule"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=h)
    _style_header(ws, 3, 3)

    rows = [
        ("Internal registration / bib / person key", "row_id", "Must be unique within the file"),
        ("WSDC ID field on profile (if verified)", "wsdc_competitor_id", "Export only real IDs; never fabricate"),
        ("Legal / ticket name", "registered_name", "PII — WSDC retention policy applies"),
        ("Preferred / badge name", "preferred_name", "Optional"),
        ("Role selected for this contest", "role", "Leader or Follower only in v0.1"),
        ("Division / level for this contest", "division", "Use Lookups list; normalize aliases before export"),
        ("Contest kind", "contest_type", "Jack & Jill / Strictly / Other"),
        ("Registration vs scratched vs competed", "entry_status", "Prefer danced if known they took the floor"),
        ("Whether WSDC ID present", "has_wsdc_id", "Y/N derived from wsdc_competitor_id"),
        ("Why ID missing", "id_absent_reason", "Only when has_wsdc_id=N"),
        ("Event catalog ID / WSDC listing ID", "Event_Header.event_id", "Preferred join key"),
        ("Official event name + year", "Event_Header.event_name_canonical + event_year", "Fallback join if no event_id"),
        ("Second role at same event", "extra Competitors row", "Do not pack two roles into one cell"),
    ]
    for r, row in enumerate(rows, start=4):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 48
    ws.freeze_panes = "A4"


def build() -> Path:
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    _write_readme(ws_readme)

    ws_lookups = wb.create_sheet("Lookups")
    ranges = _write_lookups(ws_lookups)

    ws_event = wb.create_sheet("Event_Header")
    _write_event_header(ws_event, ranges)

    ws_comp = wb.create_sheet("Competitors")
    _write_competitors(ws_comp, ranges)

    ws_map = wb.create_sheet("Vendor_Mapping")
    _write_vendor_mapping(ws_map)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
